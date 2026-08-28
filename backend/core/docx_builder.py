"""
Shared Word (.docx) report-building primitives.

Mirrors the frontend's shared print primitives
(`react-app/src/components/Shared/PrintPrimitives.tsx`) — one place for the
building blocks a formal bilingual (zh-Hant + en) report needs, so every
module's docx export composes from the same pieces instead of re-deriving
python-docx boilerplate. First consumer: NCR (`services/ncr_service.py`'s
`export_docx`). Extend here, not by copy-pasting, when the next module
(OBS/ITR/...) adds its own export.

Not merged with `services/km_service.py`'s existing HTML→docx export
(htmldocx-based) — that's a different, unrelated technique for converting
rich-text article content, not a formal fielded report. Keeping them
separate avoids coupling two unrelated concerns.
"""

import os
from io import BytesIO
from urllib.parse import quote

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from fastapi.responses import StreamingResponse

DASH = "—"

_LABEL_SHADE = "F2F2F2"   # light grey — field labels
_SECTION_SHADE = "D9D9D9"  # darker grey — section headings


_EAST_ASIAN_FONT = "標楷體"  # DFKai-SB — formal-report convention per user request 2026-08-29

def _set_font(style_obj, size=None, bold=None):
    """Calibri (Latin) + 標楷體 (East Asian)."""
    style_obj.font.name = "Calibri"
    if size:
        style_obj.font.size = size
    if bold is not None:
        style_obj.font.bold = bold
    rpr = style_obj.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), _EAST_ASIAN_FONT)


def _run_font(run, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    if size:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), _EAST_ASIAN_FONT)


def _vcenter(cell):
    """Vertically center a cell's content — python-docx cells default to
    top-aligned, which reads poorly for label/value and sign-off grids."""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return cell


def _shade_cell(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _cell_border(cell, sides=("top", "bottom", "left", "right"), sz=4, color="000000"):
    """Explicit per-cell borders — needed for elements NOT using the
    'Table Grid' style (e.g. a single-cell field box, a sign-off cell)."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in sides:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def new_document() -> Document:
    """A blank Document with the house font + A4-friendly margins set."""
    doc = Document()
    _set_font(doc.styles["Normal"], Pt(10.5))
    for hlevel in ["Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        if hlevel in doc.styles:
            _set_font(doc.styles[hlevel])
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
    return doc


def add_masthead(doc: Document, title_zh: str, title_en: str, doc_no: str,
                  rev: str = None, status: str = None,
                  company_name: str = "［ 公司名稱 Company Name ］"):
    """Header block: company placeholder + doc identity, mirrors the print
    template's `.doc-head` row. Company branding is a placeholder here too
    (BACKLOG #15 Stage A — same open item as the HTML print templates)."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Inches(4.2)
    right.width = Inches(2.3)
    _vcenter(left)
    _vcenter(right)

    p = left.paragraphs[0]
    r = p.add_run(company_name)
    _run_font(r, Pt(13), bold=True)
    p2 = left.add_paragraph()
    r2 = p2.add_run("品質管理 — Quality Management")
    _run_font(r2, Pt(9))

    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r3 = right.paragraphs[0].add_run(doc_no or "(自動 auto)")
    _run_font(r3, Pt(11), bold=True)
    meta = right.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run_font(meta.add_run(f"版次 Rev：{rev or DASH}　狀態：{(status or '').upper()}"), Pt(9))

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title_p.add_run(f"{title_zh} {title_en}")
    _run_font(tr, Pt(15), bold=True)
    doc.add_paragraph()


def add_section_heading(doc: Document, num: str, zh: str, en: str):
    """Shaded full-width heading row, mirrors `.sec-head`."""
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.rows[0].cells[0]
    _shade_cell(cell, _SECTION_SHADE)
    _vcenter(cell)
    p = cell.paragraphs[0]
    r = p.add_run(f"{num}. {zh}  {en}")
    _run_font(r, Pt(11.5), bold=True)


def add_subsection_heading(doc: Document, zh: str, en: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run(f"{zh}  {en}")
    _run_font(r, Pt(10.5), bold=True)
    r.font.underline = True


def add_field_grid(doc: Document, rows: list):
    """`rows`: list of row-tuples, each a list of (label_zh, label_en, value)
    cells (1-4 cells per row — the print template's lbl/val/lbl/val
    4-column pattern, but a row can also carry just 1-2 cells for a
    full-width value like Subject). Mirrors the repeated
    `<table><tr><td class=lbl>...<td class=val>` blocks."""
    n_cols = max(len(r) for r in rows) * 2
    table = doc.add_table(rows=0, cols=n_cols)
    table.style = "Table Grid"
    for row in rows:
        cells = table.add_row().cells
        col = 0
        span = n_cols // len(row)
        for (lbl_zh, lbl_en, value) in row:
            lbl_cell = cells[col]
            for merge_i in range(1, span - 1):
                lbl_cell = lbl_cell  # label always occupies exactly 1 col
            _shade_cell(cells[col], _LABEL_SHADE)
            _vcenter(cells[col])
            lp = cells[col].paragraphs[0]
            lr = lp.add_run(f"{lbl_zh}\n{lbl_en}")
            _run_font(lr, Pt(8.5), bold=True)
            val_start = col + 1
            val_end = col + span - 1
            if val_end > val_start:
                cells[val_start].merge(cells[val_end])
            _vcenter(cells[val_start])
            vp = cells[val_start].paragraphs[0]
            vr = vp.add_run(value if value else DASH)
            _run_font(vr, Pt(10))
            col += span
    return table


def add_field_box(doc: Document, value: str = None, guide: str = None, tall: bool = False):
    """A bordered free-text box — mirrors `.field-box` — used for long
    narrative fields (description, root cause, corrective actions, ...)."""
    table = doc.add_table(rows=1, cols=1)
    table.style = None
    cell = table.rows[0].cells[0]
    _cell_border(cell)
    _vcenter(cell)
    p = cell.paragraphs[0]
    text = value if value else (guide or DASH)
    r = p.add_run(text)
    _run_font(r, Pt(10))
    if not value:
        r.font.italic = True
        r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    if tall:
        for _ in range(2):
            cell.add_paragraph()
    return table


def add_checkbox_row(doc: Document, options: list):
    """`options`: list of (label, checked_bool). Renders inline ☑/☐."""
    p = doc.add_paragraph()
    for i, (label, checked) in enumerate(options):
        r = p.add_run(("☑ " if checked else "☐ ") + label + ("    " if i < len(options) - 1 else ""))
        _run_font(r, Pt(10), bold=checked)
    return p


def add_sign_off_grid(doc: Document, cells: list):
    """`cells`: list of dicts {num, zh, en, name, date, req}. Lays out up to
    3 signature blocks per row, mirrors `.sign-cell` — role line, blank
    signature rule, name/date row, company line."""
    per_row = 3
    for start in range(0, len(cells), per_row):
        chunk = cells[start:start + per_row]
        table = doc.add_table(rows=1, cols=len(chunk))
        table.style = "Table Grid"
        for i, spec in enumerate(chunk):
            tc = table.rows[0].cells[i]
            _cell_border(tc)
            _vcenter(tc)
            role_p = tc.paragraphs[0]
            role_text = f"{spec['num']} {spec['zh']} {spec['en']}"
            if spec.get("req"):
                role_text += f"  [{spec['req']}]"
            _run_font(role_p.add_run(role_text), Pt(9), bold=True)
            tc.add_paragraph()  # blank line for physical signature
            nd = tc.add_paragraph()
            _run_font(nd.add_run(f"姓名：{spec.get('name') or '____________'}"), Pt(8.5))
            dt = tc.add_paragraph()
            _run_font(dt.add_run(f"日期：{spec.get('date') or '____________'}"), Pt(8.5))
            co = tc.add_paragraph()
            _run_font(co.add_run("單位 Company：____________"), Pt(8.5))
        doc.add_paragraph()


def add_photo_section(doc: Document, title: str, image_paths: list):
    """`image_paths`: local filesystem paths (already resolved + existence
    checked by the caller). 2-per-row grid, mirrors `.photo-grid`."""
    p = doc.add_paragraph()
    _run_font(p.add_run(title), Pt(10.5), bold=True)
    if not image_paths:
        np = doc.add_paragraph()
        r = np.add_run("（無照片 No photos）")
        _run_font(r, Pt(9))
        r.font.italic = True
        return
    for start in range(0, len(image_paths), 2):
        pair = image_paths[start:start + 2]
        table = doc.add_table(rows=1, cols=2)
        for i, img_path in enumerate(pair):
            cell = table.rows[0].cells[i]
            _vcenter(cell)
            cp = cell.paragraphs[0]
            run = cp.add_run()
            try:
                run.add_picture(img_path, width=Inches(2.9))
            except Exception:
                _run_font(cp.add_run("（圖片載入失敗 image failed to load）"), Pt(8))


def finalize_response(doc: Document, filename_base: str) -> StreamingResponse:
    """Serialize + StreamingResponse with a browser-safe filename header —
    same convention as km_service.py's export_docx."""
    import re
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    ascii_filename = re.sub(r"[^a-zA-Z0-9_\-]", "_", filename_base)[:80] or "export"
    encoded_name = quote(f"{filename_base[:80]}.docx")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f"attachment; filename=\"{ascii_filename}.docx\"; filename*=UTF-8''{encoded_name}"
        }
    )


def resolve_local_upload_path(url_or_path: str, upload_root: str) -> str:
    """Resolve either a new-style `.../api/files/download/<relpath>` URL or a
    legacy `/uploads/<relpath>`-style string to a local filesystem path,
    with a path-traversal guard. Returns None if it can't be resolved to a
    real file under `upload_root`."""
    if not url_or_path:
        return None
    if "/api/files/download/" in url_or_path:
        rel = url_or_path.split("/api/files/download/", 1)[1]
    elif "/uploads/" in url_or_path:
        rel = url_or_path.split("/uploads/", 1)[1]
    else:
        rel = url_or_path.lstrip("/")
    rel = rel.split("?")[0]  # strip any query string (auth token, cache-bust, ...)
    full = os.path.realpath(os.path.join(upload_root, rel))
    root = os.path.realpath(upload_root)
    if not full.startswith(root):
        return None
    return full if os.path.isfile(full) else None
