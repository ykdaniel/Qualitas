"""
Meeting Minutes Service

Business logic layer for the Meeting Minutes module
"""

import json
import re
import uuid
import logging
from datetime import datetime
from typing import List, Optional

from sqlalchemy.exc import IntegrityError

import models
import schemas
from repositories.meeting_minutes_repository import MeetingMinutesRepository
from core.scope import ScopeForbidden, record_in_scope, enforce_create_scope, enforce_update_scope
from core.utils import (
    _json_serialize,
    _resolve_vendor_id,
    generate_reference_no,
    reclaim_reference_no,
    log_audit,
    WorkflowEngine
)

logger = logging.getLogger(__name__)


def _entry_kind(entry: dict) -> str:
    """Same reading as the front end's normalizeDiscussionLog (MeetingMinutes/discussionLog.ts): rows saved before
    MEETING-TABS-2026-001 have no `kind` — level 1 (or a dotted number without a level) is an item, the rest a category."""
    kind = entry.get('kind')
    if kind in ('category', 'subcategory', 'item'):
        return kind
    if entry.get('level') == 1:
        return 'item'
    if entry.get('level') is None:
        # flat rows saved 2026-08-29 … 08-31 already had an owner / status: action rows, not headings
        if 'owner' in entry or 'status' in entry:
            return 'item'
        if '.' in str(entry.get('no', '')):
            return 'item'
    return 'category'


def carry_forward_discussion_log(raw) -> Optional[str]:
    """The discussion register the next occurrence starts with (MEETING-TABS-2026-001, user decision 2026-10-11:
    only unfinished items are brought forward). A COPY, not a reference: the earlier meeting keeps its own record.

    Items still Open are kept with their Action by / Target date and every progress update, together with the
    categories / sub-categories above them; a heading with nothing open under it is not carried (user decision
    2026-10-11), and Closed items stay only in the meeting that closed them.
    Everything brought forward is marked `carried`, which the form shows grey / read-only, and numbers are not
    changed — people refer to an item by its number across meetings. Returns None when nothing is carried.
    """
    try:
        entries = json.loads(raw) if isinstance(raw, str) else raw
    except (TypeError, ValueError):
        return None
    if not isinstance(entries, list):
        return None
    rows = []
    for e in entries:
        if not isinstance(e, dict) or not str(e.get('no', '')).strip():
            continue
        kind = _entry_kind(e)
        if kind == 'item' and e.get('status') == 'Closed':
            continue
        rows.append((str(e['no']).strip(), kind, e))
    open_items = [no for no, kind, _ in rows if kind == 'item']
    carried = []
    for no, kind, e in rows:
        if kind != 'item' and not any(i.startswith(f"{no}.") for i in open_items):
            continue
        content = e.get('content')
        if content is None:
            content = ' / '.join(str(e[k]) for k in ('topic', 'discussion', 'decision') if e.get(k))
        row = {'no': no, 'kind': kind, 'level': no.count('.'), 'content': content, 'carried': True}
        if kind == 'item':
            row.update({
                'raisedDate': e.get('raisedDate') or '',
                'owner': e.get('owner') or '',
                'targetDate': e.get('targetDate') or '',
                'status': 'Open',
                'updates': [{**u, 'carried': True} for u in (e.get('updates') or [])
                            if isinstance(u, dict) and isinstance(u.get('content'), str)],
            })
        carried.append(row)
    return json.dumps(carried, ensure_ascii=False) if carried else None


class MeetingMinutesService:
    """Service layer for Meeting Minutes business logic"""

    def __init__(self, repo: MeetingMinutesRepository):
        self.repo = repo

    def get_meeting_minutes_list(self, skip: int = 0, limit: int = 500, scope=None, **filters) -> List[models.MeetingMinutes]:
        """Get list of Meeting Minutes records with optional filters"""
        return self.repo.get_all(skip, limit, scope=scope, **filters)

    def get_meeting_minutes(self, meeting_id: str, scope=None) -> Optional[models.MeetingMinutes]:
        """Get a single Meeting Minutes record by ID"""
        obj = self.repo.get_by_id(meeting_id)
        return obj if record_in_scope(obj, scope) else None

    def create_meeting_minutes(self, meeting_create: schemas.MeetingMinutesCreate,
                               user_id: int = None, username: str = None, scope=None) -> models.MeetingMinutes:
        """Create a new Meeting Minutes record"""
        try:
            data = _json_serialize(
                meeting_create.model_dump(),
                ['attendees', 'discussionLog', 'attachments']
            )

            vendor_name = data.pop('vendor', None)
            if vendor_name:
                data['vendor_id'] = _resolve_vendor_id(self.repo.db, vendor_name)

            # P0 data isolation: confine the new record to the caller's scope
            # (forces vendor_id for contractor users; validates project_id).
            enforce_create_scope(data, scope)

            if not data.get('documentNumber'):
                data['documentNumber'] = generate_reference_no(
                    self.repo.db, vendor_name or '', 'meeting'
                )
            # First occurrence of a series starts at "1.0" (BACKLOG #18) —
            # create_new_occurrence bumps this for subsequent ones.
            if not data.get('rev'):
                data['rev'] = '1.0'

            db_meeting = models.MeetingMinutes(**data)
            if not db_meeting.id:
                db_meeting.id = str(uuid.uuid4())

            created = self.repo.create(db_meeting)

            log_audit(
                self.repo.db, "CREATE", "MeetingMinutes", created.id, created.documentNumber,
                new_value=meeting_create.model_dump(), user_id=user_id, username=username
            )
            # repo.create() already committed the row itself, but that
            # happened BEFORE log_audit added its row — without a commit
            # here, the audit entry silently rolls back when the session
            # closes (confirmed empirically: audit_logs was completely
            # empty for MeetingMinutes despite extensive testing).
            self.repo.db.commit()

            return created
        except Exception as e:
            logger.error(f"Error creating Meeting Minutes: {e}", exc_info=True)
            raise e

    def create_new_occurrence(self, meeting_id: str, user_id: int = None,
                              username: str = None, scope=None) -> Optional[models.MeetingMinutes]:
        """
        Create the next occurrence of a recurring meeting series (BACKLOG
        #18) — reuses the source row's documentNumber (the whole series
        shares one document number for its lifetime) and bumps rev
        ("1.0" -> "2.0" -> ...); carries forward vendor/project/meetingType/
        organizer/location/attendees so the common case (same people, same
        format, different week) doesn't need retyping. Of the discussion
        register only the unfinished items are carried forward (with their
        categories and progress updates — carry_forward_discussion_log,
        MEETING-TABS-2026-001), and meetingDate is deliberately left blank (not defaulted to
        today) — the user fills in the actual date.

        Unlike ITR's create_reinspection / create_ncr_from_itr (which mint
        a brand-new documentNumber for a DIFFERENT module and link back via
        FK), this is deliberately a same-series, same-number derivation —
        see the unique-index migration in db_migrations.py that makes this
        legal (composite UNIQUE(documentNumber, rev) instead of a bare
        UNIQUE(documentNumber)).
        """
        try:
            source = self.repo.get_by_id(meeting_id)
            if not source or not record_in_scope(source, scope):
                return None

            siblings = self.repo.get_all_by_document_number(source.documentNumber)
            max_n = 0
            for s in siblings:
                try:
                    max_n = max(max_n, int(float(s.rev or 0)))
                except (TypeError, ValueError):
                    continue
            next_rev = f"{max_n + 1}.0"

            today = datetime.now().strftime('%Y-%m-%d')
            # Strip a trailing " - YYYY-MM-DD" this same flow would have
            # appended on a prior occurrence, so chained occurrences don't
            # accumulate multiple dates in the title over time.
            base_title = re.sub(r'\s*[-–]\s*\d{4}-\d{2}-\d{2}\s*$', '', (source.title or '').strip())
            new_title = f"{base_title} - {today}" if base_title else today

            new_meeting = models.MeetingMinutes(
                id=str(uuid.uuid4()),
                project_id=source.project_id,
                vendor_id=source.vendor_id,
                documentNumber=source.documentNumber,  # shared, not regenerated
                rev=next_rev,
                status="Draft",
                title=new_title,
                meetingType=source.meetingType,
                meetingDate=None,
                meetingTime=None,
                meetingEndTime=None,
                location=source.location,
                organizer=source.organizer,
                attendees=source.attendees,  # already a JSON string column — plain copy
                discussionLog=carry_forward_discussion_log(source.discussionLog),
                attachments=None,
                createdAt=datetime.now().isoformat(),
                updatedAt=datetime.now().isoformat(),
            )
            # No enforce_create_scope() call needed: every field above is
            # inherited straight from `source`, which record_in_scope()
            # already confirmed is within the caller's scope.

            self.repo.db.add(new_meeting)
            try:
                self.repo.db.commit()
            except IntegrityError:
                self.repo.db.rollback()
                raise ValueError(
                    "Another occurrence was just created for this series — please retry."
                )
            self.repo.db.refresh(new_meeting)

            log_audit(
                self.repo.db, "CREATE", "MeetingMinutes", new_meeting.id, new_meeting.documentNumber,
                new_value={"source": "NEW_OCCURRENCE", "sourceMeetingId": meeting_id, "rev": next_rev},
                user_id=user_id, username=username
            )
            log_audit(
                self.repo.db, "CREATE_OCCURRENCE", "MeetingMinutes", meeting_id, source.documentNumber,
                new_value={"newMeetingId": new_meeting.id, "newRev": next_rev},
                user_id=user_id, username=username
            )
            # See create_meeting_minutes' comment: log_audit only add()s,
            # nothing commits it otherwise — same fix applied here from the
            # start rather than repeating today's earlier bug.
            self.repo.db.commit()

            return new_meeting
        except ValueError as e:
            raise e
        except Exception as e:
            logger.error(f"Error creating new occurrence from Meeting Minutes {meeting_id}: {e}", exc_info=True)
            raise e

    def update_meeting_minutes(self, meeting_id: str, meeting_update: schemas.MeetingMinutesUpdate,
                               user_id: int = None, username: str = None, scope=None) -> Optional[models.MeetingMinutes]:
        """Update an existing Meeting Minutes record"""
        try:
            db_meeting = self.repo.get_by_id(meeting_id)
            if not db_meeting or not record_in_scope(db_meeting, scope):
                return None

            if meeting_update.status and not WorkflowEngine.validate_transition(
                "MeetingMinutes", db_meeting.status, meeting_update.status
            ):
                raise ValueError(
                    f"Invalid status transition from {db_meeting.status} to {meeting_update.status}"
                )

            old_val = {c.name: getattr(db_meeting, c.name) for c in db_meeting.__table__.columns}
            d = meeting_update.model_dump(exclude_unset=True)
            d = _json_serialize(d, ['attendees', 'discussionLog', 'attachments'])

            # Guard: Published and Void meeting minutes are locked — a
            # distributed/voided meeting record must not be silently
            # rewritten after the fact. The one carve-out: a Published
            # record may still transition to Void (status field only,
            # nothing else in the same request) — WorkflowEngine's
            # validate_transition above already rejects any other status
            # change from Published, so this only needs to permit that one
            # legal case. Compare against the DB's current values, not mere
            # key-presence, since the frontend resends the whole record on
            # every save.
            if db_meeting.status in ('Published', 'Void'):
                voiding = db_meeting.status == 'Published' and meeting_update.status == 'Void'
                changed_fields = {f for f in d if d[f] != old_val.get(f)}
                disallowed = changed_fields - ({'status'} if voiding else set())
                if disallowed:
                    raise ValueError(
                        f"Cannot modify meeting minute '{db_meeting.documentNumber}' "
                        f"(status: {db_meeting.status})."
                    )

            if 'vendor' in d:
                vendor_name = d.pop('vendor')
                d['vendor_id'] = _resolve_vendor_id(self.repo.db, vendor_name)

            enforce_update_scope(d, scope)

            # rev is user-editable, and (documentNumber, rev) is unique: a rev
            # already used by another occurrence of this series is a 400 for
            # the user to fix, not a 500.
            try:
                updated = self.repo.update(db_meeting, d)
            except IntegrityError:
                self.repo.db.rollback()
                # only a rev that actually changed can collide; anything else keeps its own error
                if 'rev' not in d or d['rev'] == old_val.get('rev'):
                    raise
                raise ValueError(
                    f"Rev '{d.get('rev')}' already exists for meeting minute "
                    f"'{db_meeting.documentNumber}' — choose another rev."
                )

            log_audit(
                self.repo.db, "UPDATE", "MeetingMinutes", meeting_id, updated.documentNumber,
                old_value=old_val, new_value=meeting_update.model_dump(exclude_unset=True),
                user_id=user_id, username=username
            )
            # See create_meeting_minutes' comment — log_audit needs a
            # commit after it, repo.update()'s own commit happened too
            # early to cover it.
            self.repo.db.commit()

            return updated
        except ValueError as e:
            raise e
        except Exception as e:
            logger.error(f"Error updating Meeting Minutes {meeting_id}: {e}", exc_info=True)
            raise e

    def delete_meeting_minutes(self, meeting_id: str, user_id: int = None, username: str = None, scope=None) -> bool:
        """Delete a Meeting Minutes record"""
        try:
            db_meeting = self.repo.get_by_id(meeting_id)
            if not db_meeting or not record_in_scope(db_meeting, scope):
                return False

            # Void records can always be deleted (mirrors ncr_service.py's
            # delete_ncr guard) — voiding first keeps a *published* record's
            # number visible for audit purposes instead of it silently
            # vanishing. Draft is the one exception: it was never published,
            # so nobody could have referenced its number externally — direct
            # delete is allowed, and its number is best-effort reclaimed
            # (see reclaim_reference_no) so it doesn't leave an unexplained
            # gap the way a published-then-voided record's number does.
            if db_meeting.status not in ('Void', 'Draft'):
                raise ValueError(
                    f"Cannot delete Meeting Minutes '{db_meeting.documentNumber}' "
                    f"with status '{db_meeting.status}'. Please Void it first, then delete."
                )

            is_draft = db_meeting.status == 'Draft'
            vendor_name = db_meeting.vendor or ''

            old_val = {c.name: getattr(db_meeting, c.name) for c in db_meeting.__table__.columns}
            self.repo.delete(db_meeting)

            if is_draft:
                # Must run AFTER the delete (repo.delete commits) so the
                # table scan inside reclaim_reference_no correctly excludes
                # the row that was just removed — see its docstring for why
                # this resyncs to the actual remaining max rather than a
                # naive decrement.
                reclaim_reference_no(self.repo.db, vendor_name, 'meeting')

            log_audit(
                self.repo.db, "DELETE", "MeetingMinutes", meeting_id, db_meeting.documentNumber,
                old_value=old_val, user_id=user_id, username=username
            )
            # reclaim_reference_no only flush()es (it shares the sequence
            # lock/pattern with generate_reference_no, which also only
            # flushes — callers there rely on a LATER repo.create/update
            # commit to persist it) and log_audit only db.add()s ("commit
            # is handled externally" — see its docstring). self.repo.delete()
            # above already committed the row deletion itself, but that
            # commit happened BEFORE these two calls, so nothing has
            # persisted them yet — without this, both silently roll back
            # when the request's session closes.
            self.repo.db.commit()

            return True
        except Exception as e:
            logger.error(f"Error deleting Meeting Minutes {meeting_id}: {e}", exc_info=True)
            raise e

    def export_docx(self, meeting_id: str, scope=None):
        """Word (.docx) copy of the printed minutes (MEETING-TABS-2026-001 R15, user 2026-10-11: "要補word匯出").

        Same content and order as MeetingMinutesPrintTemplate.tsx: masthead (number, rev, status), basic information
        and attendees, the discussion register (two- and three-level; items with raised date, text, progress notes in
        date order, Action by, target date and status), the attachment list, sign-off — portrait A4 — and then the
        photo appendix: the images on the Attachments tab, landscape A4, four (2 × 2) per page (user decision
        2026-10-11). Built with core/docx_builder.py like the NOI / ITR / NCR / Audit exports.

        Raises ValueError (-> 404 in the router) when the meeting does not exist or is outside the caller's scope.
        """
        from docx.enum.section import WD_ORIENT, WD_SECTION
        from docx.shared import Cm, Pt
        from core import docx_builder as db
        from core.uploads import upload_root as _upload_root

        meeting = self.get_meeting_minutes(meeting_id, scope=scope)
        if not meeting:
            raise ValueError("Meeting Minutes not found")

        def _json_list(raw):
            try:
                value = json.loads(raw) if isinstance(raw, str) else raw
            except (TypeError, ValueError):
                return []
            return value if isinstance(value, list) else []

        attendees = [{k: _text(a.get(k)) for k in ("company", "name", "role")} for a in _json_list(meeting.attendees) if isinstance(a, dict)]
        register = minutes_register(meeting.discussionLog)
        upload_root = _upload_root()
        files = self.repo.db.query(models.Attachment).filter(
            models.Attachment.entity_type == "meeting",
            models.Attachment.entity_id == meeting.id,
            models.Attachment.category == "attachment",
            models.Attachment.is_deleted == False,  # noqa: E712
        ).all()
        # upload order, as on the printed photo pages (photo numbers follow the order they were added)
        files.sort(key=lambda f: f.uploaded_at or "")
        photos = [f for f in files if _is_image(f)]
        doc_ref = f"{meeting.documentNumber or ''}{f' ({meeting.rev})' if meeting.rev else ''}"

        doc = db.new_document()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        db.add_masthead(doc, "會議記錄", "MEETING MINUTES", doc_no=meeting.documentNumber, rev=meeting.rev,
                        status=meeting.status or "Draft")

        # 1. Basic information and attendees in ONE table, the company's minutes layout (user, 2026-10-11, with a sample:
        #    "出席人員改成這樣跟基本資料合在一起") — the same structure as the print's .mom-header table
        _section_heading(doc, "1", "基本資料", "Basic Information")
        _add_basic_info_table(doc, meeting, attendees)

        # 2. Discussions / actions
        _section_heading(doc, "2", "討論事項", "Discussions / Actions")
        _add_register_table(doc, register)

        # 3. Attachments
        _section_heading(doc, "3", "附件", "Attachments")
        if files:
            for i, f in enumerate(files, start=1):
                note = f"　（照片 {photos.index(f) + 1}，見照片附錄 See photo appendix）" if f in photos else ""
                db.add_paragraph(doc, f"{i}. {f.file_name}{note}", size=9.5).paragraph_format.space_after = Pt(2)
        else:
            db.add_paragraph(doc, "（無附件 No attachments）", size=9, italic=True)

        # 4. Sign-off
        _section_heading(doc, "4", "簽核", "Sign-off")
        db.add_sign_off_grid(doc, [
            {"num": "4.1", "zh": "記錄人", "en": "Recorded By", "name": meeting.organizer, "date": format_minutes_date(meeting.meetingDate)},
            {"num": "4.2", "zh": "廠商代表", "en": "Contractor Rep.", "name": meeting.vendor},
        ])
        db.add_paragraph(doc, "會議記錄｜一經發布不可修改，如需修正請建立新版次 — Meeting Minutes: locked once Published; "
                              "create a new occurrence to amend", size=8, italic=True)

        # Photo appendix: landscape A4, 2 × 2 per page
        pages = [photos[i:i + 4] for i in range(0, len(photos), 4)]
        if pages:
            appendix = doc.add_section(WD_SECTION.NEW_PAGE)
            appendix.orientation = WD_ORIENT.LANDSCAPE
            appendix.page_width, appendix.page_height = Cm(29.7), Cm(21)
            for page_no, group in enumerate(pages):
                if page_no:
                    doc.add_page_break()
                db.add_paragraph(doc, f"照片 Photos — {doc_ref}　　{page_no + 1} / {len(pages)}", size=10.5, bold=True)
                _add_photo_grid(doc, group, page_no * 4, upload_root)

        name = f"{meeting.documentNumber or 'Meeting-Minutes'}_Rev{meeting.rev}" if meeting.rev else (meeting.documentNumber or "Meeting-Minutes")
        return db.finalize_response(doc, name)


# ── Word export helpers (R15) ──

_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_IMAGE_NAME = re.compile(r"\.(png|jpe?g|gif|webp|bmp)$", re.IGNORECASE)


def format_minutes_date(value) -> str:
    """"2026-07-21" -> "21 Jul 2026", as formatMinutesDate in discussionLog.ts; anything else as it is."""
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", str(value or ""))
    if not m or not 1 <= int(m.group(2)) <= 12:
        return str(value or "")
    return f"{int(m.group(3))} {_MONTHS[int(m.group(2)) - 1]} {m.group(1)}"


def _text(value) -> str:
    return "" if value is None else str(value)


def _no_key(no: str):
    return [int(s) if s.isdigit() else 0 for s in no.split(".")]


def minutes_register(raw) -> list:
    """The register as the print reads it (normalizeDiscussionLog): every stored shape — current two/three levels,
    legacy two-level and flat rows — in tree order; an item's progress notes in date order, same day as written."""
    try:
        entries = json.loads(raw) if isinstance(raw, str) else raw
    except (TypeError, ValueError):
        return []
    rows = []
    for e in entries if isinstance(entries, list) else []:
        if not isinstance(e, dict) or not str(e.get("no", "")).strip():
            continue
        kind = _entry_kind(e)
        content = e.get("content")
        if content is None:
            content = " / ".join(str(e[k]) for k in ("topic", "discussion", "decision") if e.get(k))
        # every value as text: the register is free JSON, a hand-written API call could put a number anywhere
        row = {"no": str(e["no"]).strip(), "kind": kind, "content": _text(content)}
        if kind == "item":
            notes = [{"date": _text(u.get("date")), "content": u["content"], "carried": bool(u.get("carried"))}
                     for u in (e.get("updates") or []) if isinstance(u, dict) and isinstance(u.get("content"), str)]
            row.update({
                "raisedDate": _text(e.get("raisedDate")), "owner": _text(e.get("owner")), "targetDate": _text(e.get("targetDate")),
                "status": "Closed" if e.get("status") == "Closed" else "Open",
                # sorted() is stable: the same day keeps the order the notes were written in
                "updates": sorted(notes, key=lambda u: u["date"]),
            })
        rows.append(row)
    return sorted(rows, key=lambda r: _no_key(r["no"]))


def _section_heading(doc, num: str, zh: str, en: str):
    """The minutes' main headings (1. 基本資料 …): bold text without the grey band docx_builder.add_section_heading
    draws (user, 2026-10-11, looking at the Word file: "大項目 不要灰底"). Only this export; the other modules keep theirs."""
    from docx.shared import Pt
    from core import docx_builder as db

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    db._run_font(p.add_run(f"{num}. {zh}  {en}"), Pt(11.5), bold=True)
    return p


def _add_basic_info_table(doc, meeting, attendees: list):
    """SUBJECT / DATE & TYPE + CONTRACTOR / VENUE & TIME + RECORDED BY / Attendees (Company / Name / Position), one
    bordered table of five columns: label | company | name | label | value; Position spans the last two."""
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    from core import docx_builder as db

    widths = [Cm(2.9), Cm(3.2), Cm(5.6), Cm(2.4), Cm(3.3)]
    table = doc.add_table(rows=0, cols=5)
    table.style = "Table Grid"
    table.autofit = False
    for i, w in enumerate(widths):
        table.columns[i].width = w

    def write(cell, lines, bold=False, size=10):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for n, line in enumerate(lines):
            p = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            db._run_font(p.add_run(line), Pt(size), bold=bold)
        return cell

    def label(cell, zh, en):
        return write(cell, [zh, en], bold=True, size=9)

    def row():
        cells = table.add_row().cells
        for i, c in enumerate(cells):
            c.width = widths[i]
            # tight even where nothing is written (the rows under the merged ATTENDEES label)
            c.paragraphs[0].paragraph_format.space_before = Pt(1)
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
        return cells

    time_range = " ~ ".join(t for t in (meeting.meetingTime, meeting.meetingEndTime) if t)
    date_type = "　".join(x for x in (format_minutes_date(meeting.meetingDate), meeting.meetingType) if x)
    c = row()
    label(c[0], "主題", "SUBJECT")
    write(c[1].merge(c[4]), [meeting.title or db.DASH])
    c = row()
    label(c[0], "日期／類型", "DATE / TYPE")
    write(c[1].merge(c[2]), [date_type or db.DASH])
    label(c[3], "廠商", "CONTRACTOR")
    write(c[4], [meeting.vendor or db.DASH])
    c = row()
    label(c[0], "地點與時間", "VENUE & TIME")
    write(c[1].merge(c[2]), [x for x in (meeting.location, time_range) if x] or [db.DASH])
    label(c[3], "記錄人", "RECORDED BY")
    write(c[4], [meeting.organizer or db.DASH])
    head = row()
    write(head[1], ["公司 Company"], bold=True, size=9)
    write(head[2], ["姓名 Name"], bold=True, size=9)
    write(head[3].merge(head[4]), ["職稱 Position"], bold=True, size=9)
    last = head
    if attendees:
        for a in attendees:
            last = row()
            write(last[1], [a["company"] or db.DASH])
            write(last[2], [a["name"] or db.DASH])
            write(last[3].merge(last[4]), [a["role"] or db.DASH])
    else:
        last = row()
        write(last[1].merge(last[4]), ["（無資料 No attendees）"], size=9)
    label(head[0].merge(last[0]), "出席人員", "ATTENDEES")
    return table


def _is_image(f) -> bool:
    return (f.mime_type or "").startswith("image/") or bool(_IMAGE_NAME.search(f.file_name or ""))


def _add_register_table(doc, register: list):
    """# / DISCUSSIONS / ACTIONS / Action by / Target Date & Status — widths as the print (8 / 62 / 14 / 16 %)."""
    from docx.shared import Cm, Pt, RGBColor
    from core import docx_builder as db

    widths = [Cm(1.4), Cm(10.8), Cm(2.4), Cm(2.8)]
    grey = RGBColor(0x80, 0x80, 0x80)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    for i, w in enumerate(widths):          # the grid too, not only each cell: some readers size columns from it
        table.columns[i].width = w

    def text(cell, value, size=9, bold=False, color=None, indent=False, first=False):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        if indent:
            p.paragraph_format.left_indent = Cm(0.35)
        db._run_font(p.add_run(value), Pt(size), bold=bold, color=color)
        return p

    for i, h in enumerate(["#", "DISCUSSIONS / ACTIONS", "Action by", "Target Date / Status"]):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        db._shade_cell(cell, "F2F2F2")
        text(cell, h, bold=True, first=True)
    if not register:
        cells = table.add_row().cells
        cells[0].merge(cells[3])
        text(cells[0], "（無資料 No discussion items）", color=grey, first=True)
        return table
    for e in register:
        cells = table.add_row().cells
        for i, cell in enumerate(cells):
            cell.width = widths[i]
        if e["kind"] != "item":
            db._shade_cell(cells[0], "D9D9D9" if e["kind"] == "category" else "EEEEEE")
            label = f"{e['no']}.0" if e["kind"] == "category" else e["no"]
            text(cells[0], label, bold=True, first=True)
            merged = cells[1].merge(cells[3])
            db._shade_cell(merged, "D9D9D9" if e["kind"] == "category" else "EEEEEE")
            text(merged, e["content"], bold=True, first=True)
            continue
        text(cells[0], e["no"], first=True)
        body = cells[1]
        first = True
        if e["raisedDate"]:
            text(body, format_minutes_date(e["raisedDate"]), size=8, color=grey, first=True)
            first = False
        for line in (e["content"] or db.DASH).split("\n"):
            text(body, line, indent=bool(e["raisedDate"]), first=first)
            first = False
        for u in e["updates"]:
            color = grey if u.get("carried") else None
            if u.get("date"):
                text(body, format_minutes_date(u["date"]), size=8, color=grey)
            for line in u["content"].split("\n"):
                text(body, line, indent=True, color=color)
        text(cells[2], e["owner"] or db.DASH, first=True)
        if e["targetDate"]:
            text(cells[3], format_minutes_date(e["targetDate"]), first=True)
            text(cells[3], e["status"])
        else:
            text(cells[3], e["status"], first=True)
    return table


def _add_photo_grid(doc, group: list, offset: int, upload_root: str):
    """Up to four photos, 2 × 2, each fitted into its box (12.4 × 7.0 cm) with its number and file name below."""
    from docx.image.image import Image as DocxImage
    from docx.shared import Cm, Pt
    from core import docx_builder as db

    box_w, box_h = Cm(12.4), Cm(7.0)     # two rows of photo + caption fit the 18 cm landscape page with room to spare
    table = doc.add_table(rows=(len(group) + 1) // 2, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for col in table.columns:
        col.width = Cm(12.9)
    for i, f in enumerate(group):
        cell = table.rows[i // 2].cells[i % 2]
        cell.width = Cm(12.9)
        p = cell.paragraphs[0]
        path = db.resolve_local_upload_path(f.file_path, upload_root)
        try:
            if not path:
                raise ValueError("missing")
            img = DocxImage.from_file(path)
            scale = min(box_w / img.width, box_h / img.height)  # EMU sizes: fit inside the box, keep the proportions
            p.add_run().add_picture(path, width=int(img.width * scale), height=int(img.height * scale))
        except Exception:
            db._run_font(p.add_run("（圖片無法顯示 image could not be shown）"), Pt(9))
        cap = cell.add_paragraph()
        db._run_font(cap.add_run(f"照片 Photo {offset + i + 1}　{f.file_name}"), Pt(8.5))
