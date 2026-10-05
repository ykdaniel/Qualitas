"""ITR-EXPORT-DOCX-2026-001: real HTTP + python-docx content verification for
GET /api/itr/{id}/export-docx — not a seed script, read-only HTTP calls against an already
running isolated stack (see isolated_stack.py). Downloads the .docx and parses it back with
python-docx to assert on actual document content, not just "the request returned 200".

Usage:
    python scripts/verification/verify_itr_export_docx.py <base_url> <username> <password> <itr_id> <out_path>
"""
import sys
import requests
from docx import Document

base_url, username, password, itr_id, out_path = sys.argv[1:6]

total = 0
failures = 0


def check(cond, msg):
    global total, failures
    total += 1
    status = "PASS" if cond else "FAIL"
    if not cond:
        failures += 1
    print(f"{status}: {msg}")


session = requests.Session()
login_resp = session.post(f"{base_url}/api/auth/login", data={"username": username, "password": password})
check(login_resp.status_code == 200, f"login succeeded (status {login_resp.status_code})")

resp = session.get(f"{base_url}/api/itr/{itr_id}/export-docx")
check(resp.status_code == 200, f"GET /api/itr/{itr_id}/export-docx returned 200 (got {resp.status_code})")
check(resp.headers.get("content-type", "").startswith("application/vnd.openxmlformats"), f"content-type is a .docx mime type (got {resp.headers.get('content-type')})")

with open(out_path, "wb") as f:
    f.write(resp.content)

doc = Document(out_path)
full_text = "\n".join(p.text for p in doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            full_text += "\n" + cell.text

check("QTS-IUX2-ITR-000002" in full_text, "the ITR's own document number appears in the .docx")
check("QTS-IUX2-NOI-000001" in full_text, "the linked NOI number appears")
check("QTS-IUX2-ITP-000001" in full_text, "the live-derived Related ITP appears (not blank)")
check("Meridian Concrete Works" in full_text, "the contractor name appears")

# Find the Checklist data table specifically (headers row contains "Item" and "Result").
checklist_table = None
for table in doc.tables:
    header_texts = [c.text for c in table.rows[0].cells]
    if any("Item" in h for h in header_texts) and any("Result" in h for h in header_texts):
        checklist_table = table
        break
check(checklist_table is not None, "a Checklist Item/Criteria/Situation/Result table exists in the document")

if checklist_table is not None:
    data_row = checklist_table.rows[1]
    situation_cell = data_row.cells[3]
    result_cell = data_row.cells[4]
    situation_paragraph_count = len(situation_cell.paragraphs)
    check(situation_paragraph_count >= 3, f"the Situation cell has multiple paragraphs (real line breaks), not one run-on line (found {situation_paragraph_count} paragraphs)")
    check("SITU-MARK-TWO-xyz789" in situation_cell.text, f"the Situation cell's full text includes the trailing marker (text: {situation_cell.text[:80]!r}...)")
    check("Pass" in result_cell.text, f"the Result cell shows a human-readable 'Pass' label (got {result_cell.text!r})")

print(f"\n=== SUMMARY: {total} checks executed, {total - failures} PASS, {failures} FAIL ===")
sys.exit(1 if failures > 0 else 0)
