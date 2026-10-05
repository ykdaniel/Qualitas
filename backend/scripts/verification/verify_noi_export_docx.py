"""NOI-EXPORT-DOCX-2026-001: real HTTP + python-docx content verification for
GET /api/noi/{id}/export-docx — not a seed script, read-only HTTP calls against an already
running isolated stack (see isolated_stack.py).

Usage:
    python scripts/verification/verify_noi_export_docx.py <base_url> <username> <password> <noi_id> <out_path> [--expect-no-attachments]
"""
import sys
import zipfile
import requests
from docx import Document

base_url, username, password, noi_id, out_path = sys.argv[1:6]
expect_no_attachments = "--expect-no-attachments" in sys.argv[6:]

total = 0
failures = 0


def check(cond, msg):
    global total, failures
    total += 1
    status = "PASS" if cond else "FAIL"
    if not cond:
        failures += 1
    print(f"{status}: {msg}")


session = requests.Session()
login_resp = session.post(f"{base_url}/api/auth/login", data={"username": username, "password": password})
check(login_resp.status_code == 200, f"login succeeded (status {login_resp.status_code})")

resp = session.get(f"{base_url}/api/noi/{noi_id}/export-docx")
check(resp.status_code == 200, f"GET /api/noi/{noi_id}/export-docx returned 200 (got {resp.status_code})")
check(resp.headers.get("content-type", "").startswith("application/vnd.openxmlformats"), f"content-type is a .docx mime type (got {resp.headers.get('content-type')})")

with open(out_path, "wb") as f:
    f.write(resp.content)

doc = Document(out_path)
full_text = "\n".join(p.text for p in doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            full_text += "\n" + cell.text

if expect_no_attachments:
    check("Attachments" not in full_text, "no 'Attachments' section heading appears when there are no attachments")
else:
    check("Attachments" in full_text, "the 'Attachments' section heading appears")
    check("seed-noi-attachment.png" in full_text, "the attachment's actual filename is listed")
    media_files = [n for n in zipfile.ZipFile(out_path).namelist() if n.startswith("word/media/")]
    check(len(media_files) == 0, f"no images are embedded (NOI attachments are listed by filename, not embedded as photos) — found media: {media_files}")

print(f"\n=== SUMMARY: {total} checks executed, {total - failures} PASS, {failures} FAIL ===")
sys.exit(1 if failures > 0 else 0)
