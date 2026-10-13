"""Audit Word export (AUDIT-EXPORT-DOCX-2026-001): GET /api/audit/{id}/export-docx returns the whole audit report.

Real login/routes on a disposable file DB (fixture shared with test_audit_hardening_http); the .docx is opened with
python-docx and its text checked. No main import, no development database access.
"""
from io import BytesIO
from urllib.parse import unquote

from docx import Document

import models
from core.security import get_password_hash
from project_membership import TEST_PROJECT_ID, bind
from test_audit_hardening_http import create, env, insert_raw, login_admin  # noqa: F401  (env is a fixture)

PW = 'Export-Test-1234'  # isolated, test-only
DOCX_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def export(env, audit_id):
    return env.client.get(f'/api/audit/{audit_id}/export-docx')


def docx_text(response):
    """Every paragraph of the document, body and table cells, in order."""
    doc = Document(BytesIO(response.content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return '\n'.join(parts)


def item_rows(response):
    """The checklist detail table (header row starts with '#'), as lists of cell texts."""
    doc = Document(BytesIO(response.content))
    for table in doc.tables:
        if table.rows[0].cells[0].text == '#':
            return [[c.text for c in row.cells] for row in table.rows[1:]]
    return None


def add_user(env, username, codes, vendor_id=None, projects=(TEST_PROJECT_ID,)):
    """A test account bound to `projects` (default: the test project the harness's audits live in; () = bind none here)."""
    with env.Session() as db:
        role = models.Role(name=f'role-{username}')
        role.permissions_rel = [db.query(models.Permission).filter_by(code=c).first() or models.Permission(code=c, description=c)
                                for c in codes]
        db.add(role)
        db.flush()
        user = models.User(username=username, email=f'{username}@example.com', is_active=True, role_id=role.id,
                           vendor_id=vendor_id, hashed_password=get_password_hash(PW))
        db.add(user)
        db.flush()
        if projects:
            bind(db, user, *projects)
        db.commit()


def login(env, username):
    response = env.client.post('/api/auth/login', data={'username': username, 'password': PW})
    assert response.status_code == 200, response.text


# results (AUDIT-GRADES-2026-001): pass / positive / major / minor; an older 'fail' is "不符合（未分級）" (graded);
# 'pending' and no status are 未評 (not graded)
ITEMS = [
    {'id': 1, 'no': '7.1', 'clause': 'Resources', 'task': 'How are resources provided?', 'status': 'pass',
     'note': 'Budget approved', 'updatedAt': '2026-10-02'},
    {'id': 2, 'no': '8.4', 'clause': 'External providers', 'task': 'How are suppliers controlled?', 'status': 'major',
     'note': 'No evaluation\nfor two vendors', 'updatedAt': '2026-10-02'},
    {'id': 3, 'no': '', 'clause': '', 'task': 'Custom question', 'status': 'pending'},
    {'id': 4, 'no': '9.2', 'clause': 'Internal audit', 'task': 'Planned intervals?'},  # no status yet → not graded
    {'id': 5, 'no': '7.2', 'clause': 'Competence', 'task': 'Training records?', 'status': 'positive'},
    {'id': 6, 'no': '7.5', 'clause': 'Documented information', 'task': 'Document control?', 'status': 'minor'},
    {'id': 7, 'no': '8.7', 'clause': 'Nonconforming outputs', 'task': 'Segregation?', 'status': 'fail'},  # older audit
]


def test_export_contains_the_whole_report(env):
    with env.Session() as db:
        project = db.get(models.Project, TEST_PROJECT_ID)  # the auditor's working project, labelled like a real one
        project.name, project.code = 'Harbour Line', 'HL'
        db.commit()
    created = create(env, title='Q4 supplier audit', project_id=TEST_PROJECT_ID, project_name='Harbour Line', location='Yard 3',
                     project_director='Director Chen', tech_lead='Tech Lin', auditor='Lead Wang', support_auditors='Support Lee',
                     audit_criteria='Contract + ISO 9001', scope_description='Receiving and storage',
                     selected_templates=['ISO 9001:2015', '供應商評鑑表'], custom_check_items=ITEMS,
                     findings='Two suppliers lack evaluation.\nFollow-up NCR to be raised.', status='In Progress').json()
    response = export(env, created['id'])
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == DOCX_TYPE
    disposition = response.headers['content-disposition']
    assert disposition.startswith('attachment;')
    assert unquote(disposition.split("filename*=UTF-8''")[1]) == f"{created['auditNo']}.docx"

    text = docx_text(response)
    for expected in ['內部稽核報告 INTERNAL AUDIT REPORT', created['auditNo'], '狀態：IN PROGRESS', '進行中 In Progress',
                     'Q4 supplier audit', '[HL] Harbour Line', 'Audit vendor', '2026-10-01', '2026-10-02', 'Yard 3',
                     'Director Chen', 'Tech Lin', 'Lead Wang', 'Support Lee', 'Contract + ISO 9001', 'ISO 9001:2015、供應商評鑑表',
                     'Receiving and storage', 'Two suppliers lack evaluation.', 'Follow-up NCR to be raised.',
                     '主任稽核員 Lead Auditor', '受稽核方代表 Auditee Representative', '專案經理 Project Director']:
        assert expected in text, expected
    # summary: 7 items, one of each grade + an older ungraded fail, 2 not graded → progress 5 / 7 = 71%
    doc = Document(BytesIO(response.content))
    summary = next(t for t in doc.tables if 'Total Items' in t.rows[0].cells[0].text)
    values = {}
    for row in summary.rows:
        texts = [c.text for c in row.cells]
        texts = [x for i, x in enumerate(texts) if i == 0 or x != texts[i - 1]]  # a merged value cell repeats its text
        values.update({texts[i].split('\n')[1]: texts[i + 1] for i in range(0, len(texts), 2)})
    assert values == {'Total Items': '7', 'Progress': '71%', 'Pass': '1', 'Positive': '1', 'Major': '1', 'Minor': '1',
                      'Not graded': '2', 'Fail (ungraded)': '1'}

    rows = item_rows(response)
    assert [r[0] for r in rows] == ['1', '2', '3', '4', '5', '6', '7']  # stored order
    # clause text and audit question in separate columns (AUDIT-GRADES-2026-001)
    assert rows[0][1:] == ['7.1', 'Resources', 'How are resources provided?', '符合 Pass', 'Budget approved', '2026-10-02']
    assert rows[1][4] == '主要不符合 Major' and rows[1][5] == 'No evaluation\nfor two vendors'  # multi-line note kept
    assert rows[2][1:5] == ['—', '—', 'Custom question', '未評 Not graded']
    assert rows[3][4] == '未評 Not graded' and rows[3][5] == '—' and rows[3][6] == '—'
    assert [r[4] for r in rows[4:]] == ['正面發現 Positive', '次要不符合 Minor', '不符合（未分級） Fail']
    table = next(t for t in Document(BytesIO(response.content)).tables if t.rows[0].cells[0].text == '#')
    assert [c.text.split('\n')[0] for c in table.rows[0].cells] == ['#', '條文編號', '條文內容', '查檢重點', '結果', '備註', '更新日期']
    assert [round(c.width.cm, 1) for c in table.rows[1].cells] == [0.8, 1.4, 4.4, 3.8, 2.2, 3.0, 1.8]


def test_export_of_an_audit_without_items_or_optional_fields(env):
    created = create(env, custom_check_items=[], selected_templates=[], location='', findings='').json()
    response = export(env, created['id'])
    assert response.status_code == 200, response.text
    text = docx_text(response)
    assert '（尚無查檢項目 No audit items）' in text
    assert item_rows(response) is None
    assert '記錄稽核發現、不符合事項與總結' in text  # empty findings → the guide text, like the NCR report


def test_legacy_rows_export_without_failing(env):
    # project name only (no project_id), list fields that are not valid JSON / not lists / hold non-objects.
    # A row without project_id is visible only to a system administrator (REF-PROJECT-SEQ B1), so export as one.
    login_admin(env)
    for templates, items in [('not json', '{"a": 1}'), ('{"x": 1}', '[1, "two", null]'), (None, None)]:
        audit_id = insert_raw(env, project_id=None, project_name='Old Project', selected_templates=templates, custom_check_items=items)
        response = export(env, audit_id)
        assert response.status_code == 200, response.text
        text = docx_text(response)
        assert 'Old Project' in text
        assert '（尚無查檢項目 No audit items）' in text


def test_a_status_that_is_not_a_string_counts_as_not_graded(env):
    # the create API accepts any JSON in custom_check_items, so a list / object / number status can be stored (R1 review)
    items = [{'id': 1, 'task': 'list', 'status': ['pass']}, {'id': 2, 'task': 'object', 'status': {'v': 'fail'}},
             {'id': 3, 'task': 'number', 'status': 1}, {'id': 4, 'task': 'ok', 'status': 'pass'}]
    created = create(env, custom_check_items=items)
    assert created.status_code == 200, created.text
    response = export(env, created.json()['id'])
    assert response.status_code == 200, response.text
    assert [r[4] for r in item_rows(response)] == ['未評 Not graded'] * 3 + ['符合 Pass']
    assert '25%' in docx_text(response)  # 1 of 4 graded, the other three not graded
    assert 'Fail (ungraded)' not in docx_text(response)  # the ungraded-fail line only shows when there is one


def summary_values(response):
    doc = Document(BytesIO(response.content))
    summary = next(t for t in doc.tables if 'Total Items' in t.rows[0].cells[0].text)
    values = {}
    for row in summary.rows:
        texts = [c.text for c in row.cells]
        texts = [x for i, x in enumerate(texts) if i == 0 or x != texts[i - 1]]
        values.update({texts[i].split('\n')[1]: texts[i + 1] for i in range(0, len(texts), 2)})
    return values


def test_progress_rounds_half_up_like_the_wizard(env):
    # 1 of 8 done = 12.5 % -> 13 % (round() would give 12); 23 of 40 = 57.5 % -> 58 % (in JavaScript 23 / 40 * 100 is
    # 57.4999..., so Math.round gave 57 — the wizard and the printouts now use the same integer formula)
    for total, done, expected in [(8, 1, '13%'), (40, 23, '58%'), (3, 1, '33%'), (3, 2, '67%'), (6, 6, '100%')]:
        items = [{'id': n, 'task': f'q{n}', 'status': 'pass' if n < done else None} for n in range(total)]
        created = create(env, custom_check_items=items).json()
        assert summary_values(export(env, created['id']))['Progress'] == expected, (total, done)


def test_characters_xml_cannot_carry_are_dropped_not_an_error(env):
    audit_id = insert_raw(env, title='Ti\x01tle', findings='Line\x1fone\nLine two',
                          custom_check_items='[{"id": 1, "no": "7.1", "clause": "Res\\u000bources", "task": "Q\\u0000?",'
                                             ' "status": "pass", "note": "bad\\ud800char"}]')
    response = export(env, audit_id)
    assert response.status_code == 200, response.text
    text = docx_text(response)
    assert 'Title' in text and 'Lineone' in text and 'Line two' in text
    assert item_rows(response)[0][2:6] == ['Resources', 'Q?', '符合 Pass', 'badchar']


def test_unknown_audit_is_404(env):
    response = export(env, 'no-such-audit')
    assert response.status_code == 404
    assert response.json() == {'detail': 'Audit not found'}


def test_contractor_scoped_user_exports_only_its_own_audits(env):
    with env.Session() as db:
        db.add(models.Contractor(id='alpha', name='Alpha Build', abbreviation='AB', status='active'))
        db.commit()
    other = create(env).json()  # vendor 'Audit vendor', created by the unrestricted auditor
    add_user(env, 'alpha-viewer', ['audit:view:all', 'audit:create:all'], vendor_id='alpha')
    login(env, 'alpha-viewer')
    own = create(env).json()
    assert own['vendor_id'] == 'alpha'
    assert export(env, own['id']).status_code == 200
    response = export(env, other['id'])
    assert response.status_code == 404  # outside the scope: same answer as a missing audit, no existence leak
    assert response.json() == {'detail': 'Audit not found'}


def test_project_scoped_user_cannot_export_another_projects_audit(env):
    with env.Session() as db:
        db.add_all([models.Project(id='p1', name='Project One'), models.Project(id='p2', name='Project Two')])
        db.commit()
    other = insert_raw(env, project_id='p2')  # an audit of project p2 (the auditor works in the test project only)
    add_user(env, 'p1-viewer', ['audit:view:all'], projects=())
    with env.Session() as db:
        user = db.query(models.User).filter_by(username='p1-viewer').one()
        db.add(models.UserProject(user_id=user.id, project_id='p1'))
        db.commit()
    login(env, 'p1-viewer')
    assert export(env, other).status_code == 404


def test_without_audit_view_or_signed_out_is_refused(env):
    created = create(env).json()
    add_user(env, 'no-audit', ['ncr:view:all'])
    login(env, 'no-audit')
    assert export(env, created['id']).status_code == 403
    env.client.cookies.clear()
    assert export(env, created['id']).status_code == 401


# ── AUDIT-BUTTONS-2026-001: the plan and the on-site checklist as Word, like the step printouts ───────────────────────
def section_headings(response):
    doc = Document(BytesIO(response.content))
    return [t.rows[0].cells[0].text for t in doc.tables if len(t.columns) == 1 and t.rows[0].cells[0].text[:2].rstrip('.').isdigit()]


def test_kind_plan_is_sections_1_to_4_plus_a_plan_sign_off(env):
    created = create(env, title='Plan export', agenda=[{'date': '2026-10-01', 'start': '09:00', 'end': '09:30', 'item': 'Opening'}],
                     custom_check_items=ITEMS).json()
    response = env.client.get(f"/api/audit/{created['id']}/export-docx?kind=plan")
    assert response.status_code == 200, response.text
    assert unquote(response.headers['content-disposition'].split("filename*=UTF-8''")[1]) == f"{created['auditNo']}-Plan.docx"
    text = docx_text(response)
    assert '內部稽核計畫 INTERNAL AUDIT PLAN' in text and 'Plan export' in text and 'Opening' in text
    assert section_headings(response) == ['1. 稽核資訊  Audit Information', '2. 稽核團隊  Audit Team', '3. 稽核範圍與準則  Scope & Criteria',
                                          '4. 稽核行程  Agenda', '5. 計畫簽核  Plan Sign-off']
    assert '查檢結果' not in text and 'Budget approved' not in text  # no results in a plan
    assert '5.1 主任稽核員 Lead Auditor' in text and '5.2 專案經理 Project Director' in text


def test_kind_checklist_has_blank_result_boxes_and_no_recorded_results(env):
    created = create(env, custom_check_items=ITEMS, location='Yard 3').json()
    response = env.client.get(f"/api/audit/{created['id']}/export-docx?kind=checklist")
    assert response.status_code == 200, response.text
    assert unquote(response.headers['content-disposition'].split("filename*=UTF-8''")[1]) == f"{created['auditNo']}-Checklist.docx"
    text = docx_text(response)
    assert '稽核查檢表 AUDIT CHECKLIST' in text and 'Yard 3' in text and '2026-10-01 ~ 2026-10-02' in text
    assert section_headings(response) == ['1. 稽核資訊  Audit Information', '2. 查檢項目  Audit Items', '3. 簽核  Sign-off']
    rows = item_rows(response)
    assert [r[0] for r in rows] == ['1', '2', '3', '4', '5', '6', '7']
    assert rows[0][1:] == ['7.1', 'Resources', 'How are resources provided?',
                           '☐ 符合 Pass\n☐ 正面發現 Positive\n☐ 主要不符合 Major\n☐ 次要不符合 Minor', '']
    # recorded results / notes are not printed
    assert 'Budget approved' not in text and '主要不符合 Major' not in text.replace('☐ 主要不符合 Major', '')
    table = next(t for t in Document(BytesIO(response.content)).tables if t.rows[0].cells[0].text == '#')
    assert [round(c.width.cm, 1) for c in table.rows[1].cells] == [0.8, 1.4, 5.0, 4.4, 2.8, 3.0]
    empty = create(env, custom_check_items=[]).json()
    assert '（尚無查檢項目 No audit items）' in docx_text(env.client.get(f"/api/audit/{empty['id']}/export-docx?kind=checklist"))


def test_kind_report_is_the_default_and_unknown_kinds_are_refused(env):
    created = create(env, custom_check_items=ITEMS).json()
    default = env.client.get(f"/api/audit/{created['id']}/export-docx")
    explicit = env.client.get(f"/api/audit/{created['id']}/export-docx?kind=report")
    assert default.status_code == explicit.status_code == 200
    assert section_headings(default) == section_headings(explicit)
    assert section_headings(default)[-1] == '7. 簽核  Sign-off'
    for bad in ['summary', 'REPORT', '', '../plan']:
        response = env.client.get(f"/api/audit/{created['id']}/export-docx", params={'kind': bad})
        assert response.status_code == 400, (bad, response.text)
        assert 'Unknown export kind' in response.json()['detail']
    # scope / not-found still answer before the kind matters for an existing kind
    assert env.client.get('/api/audit/no-such/export-docx?kind=plan').status_code == 404
