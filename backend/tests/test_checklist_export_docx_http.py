"""Checklist Word export (CHECKLIST-BACK-LABEL-2026-001 R2, user 2026-10-12: "除了列印還要有匯出word").

GET /api/checklist/{id}/export-docx returns the printed Checklist as .docx — the same sections as
ChecklistPrintTemplate.tsx, for templates and instances alike (user's choice "B. 完全照列印版": a template's file is the
blank field form — empty Actual Situation, Result "未填 Not filled" as the print view shows it, and the signature blocks). Needs checklist:view and the
record in the caller's scope (another project's checklist is a 404, like reading it). Throwaway file-backed SQLite,
real logins, real HTTP.
"""
import json
from io import BytesIO
from urllib.parse import unquote

import pytest
from docx import Document
from fastapi import FastAPI
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import models
from core import perms
from core.security import get_password_hash
from test_itr_revoke_approval_acceptance import PW, Env, _get_or_create_perm

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEMPLATE, INSTANCE, OTHER_PROJECT = "chk-template", "chk-instance", "chk-other"


@pytest.fixture
def chk_env(tmp_path):
    from database import Base, get_db
    from routers import auth as auth_router, checklist as checklist_router
    import database as database_module
    import core.security as security_module   # binds SessionLocal at import: swap it too (token blacklist)
    engine = create_engine(f"sqlite:///{tmp_path / 'checklist.db'}", connect_args={"check_same_thread": False, "timeout": 20})
    Base.metadata.create_all(engine)
    S = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    app = FastAPI()
    for r in (auth_router.router, checklist_router.router):
        app.include_router(r, prefix="/api")

    def _db():
        s = S()
        try:
            yield s
        finally:
            s.close()
    app.dependency_overrides[get_db] = _db
    orig = (database_module.engine, database_module.SessionLocal, security_module.SessionLocal)
    database_module.engine, database_module.SessionLocal, security_module.SessionLocal = engine, S, S
    try:
        _seed(S)
        yield Env(app, S)
    finally:
        database_module.engine, database_module.SessionLocal, security_module.SessionLocal = orig
        Base.metadata.drop_all(engine)
        engine.dispose()


def _user(db, username, codes, project_id=None):
    role = models.Role(name=f"role-{username}")
    role.permissions_rel = [_get_or_create_perm(db, c) for c in codes]
    db.add(role)
    db.flush()
    user = models.User(username=username, email=f"{username}@example.com", is_active=True, role_id=role.id,
                       hashed_password=get_password_hash(PW))
    db.add(user)
    db.flush()
    if project_id:
        db.add(models.UserProject(user_id=user.id, project_id=project_id))


def _seed(S):
    db = S()
    try:
        db.add(models.Project(id="EXP-P1", name="Export Review", code="EXP1"))
        db.add(models.Project(id="EXP-P2", name="Other Project", code="EXP2"))
        db.add(models.Contractor(id="EXP-V1", name="Hailong Construction", abbreviation="HL"))
        db.flush()
        _user(db, "chk_viewer", [perms.CHECKLIST_VIEW], project_id="EXP-P1")       # scoped to project 1
        _user(db, "chk_noperm", [perms.NOI_VIEW], project_id="EXP-P1")             # no checklist permission
        header = {"projectTitle": "Hai Long", "referenceNo": "QTS-EXP1-HL-ITR-000001", "location": "Grid 3-5",
                  "revision": 2, "remarks": "1. Mark 'O' for passed items.\n2. Quantify the criteria."}
        template_items = [
            {"id": 1, "item": "Material check", "criteria": "Manufacturer, size and grade per submittal",
             "situation": "", "result": ""},
            {"id": 2, "item": "Welding visual", "criteria": "No cracks\nWeld size ≥ drawing", "situation": "", "result": ""},
        ]
        db.add(models.Checklist(id=TEMPLATE, recordsNo="QTS-EXP1-HL-CHK-000001", activity="Steel erection",
                                date="2026-10-12", status="Ongoing", project_id="EXP-P1", contractor_id="EXP-V1",
                                detail_data=json.dumps({**header, "inspectionDate": "", "agreementChecked": True,
                                                        "items": template_items}, ensure_ascii=False)))
        instance_items = [
            {"id": 1, "item": "Material check", "criteria": "Per submittal", "situation": "Mill certs OK", "result": "O"},
            {"id": 2, "item": "Welding visual", "criteria": "No cracks", "situation": "Porosity at J12", "result": "X"},
            {"id": 3, "item": "Coating", "criteria": "DFT 250 µm", "situation": "", "result": "/",
             "naReason": "Coating by others"},
            {"id": 4, "item": "Legacy N/A", "criteria": "-", "situation": "", "result": "/"},
            {"id": 5, "item": "Odd code", "criteria": "-", "situation": "", "result": "??"},
        ]
        db.add(models.Checklist(id=INSTANCE, recordsNo="QTS-EXP1-HL-CHK-000002", activity="Steel erection",
                                date="2026-10-12", status="Fail", passCount=1, failCount=1, project_id="EXP-P1",
                                contractor_id="EXP-V1", noiNumber="QTS-EXP1-HL-NOI-000001", template_id=TEMPLATE,
                                detail_data=json.dumps({**header, "inspectionDate": "2026-10-12", "agreementChecked": False,
                                                        "ncrNo": "NCR-007", "reInspectionDate": "2026-10-20",
                                                        "items": instance_items}, ensure_ascii=False)))
        db.add(models.Checklist(id=OTHER_PROJECT, recordsNo="QTS-EXP2-HL-CHK-000001", activity="Elsewhere",
                                date="2026-10-12", status="Ongoing", project_id="EXP-P2", contractor_id="EXP-V1",
                                detail_data=json.dumps({"items": []})))
        db.commit()
    finally:
        db.close()


def _texts(response):
    """Body paragraphs, then every table cell's paragraphs (merged cells once)."""
    doc = Document(BytesIO(response.content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def _items_rows(response):
    doc = Document(BytesIO(response.content))
    for table in doc.tables:
        if [c.text for c in table.rows[0].cells] == ["#", "Inspection Item", "Criteria", "Actual Situation", "Result"]:
            return [[c.text for c in row.cells] for row in table.rows[1:]]
    return None


def test_template_exports_the_printed_blank_form(chk_env):
    r = chk_env.login("chk_viewer").get(f"/api/checklist/{TEMPLATE}/export-docx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == DOCX_TYPE
    assert "QTS-EXP1-HL-CHK-000001.docx" in unquote(r.headers["content-disposition"])
    text = _texts(r)
    for s in ["現場檢驗查核表 FIELD INSPECTION CHECKLIST", "QTS-EXP1-HL-CHK-000001", "版次 Rev：2", "Steel erection",
              "Hai Long", "QTS-EXP1-HL-ITR-000001", "Hailong Construction", "Grid 3-5",
              "☑ All inspection has been done", "☐ Unfinished improvement", "1. Mark 'O' for passed items.",
              "Site Engineer", "Construction Leader", "Subcontractor Representative"]:
        assert s in text, s
    rows = _items_rows(r)
    # the same five columns as the print view, situation / result blank for a template, then a blank row and -END-
    assert rows[0] == ["1", "Material check", "Manufacturer, size and grade per submittal", "", "未填 Not filled"]
    assert rows[1][2] == "No cracks\nWeld size ≥ drawing"           # an old line break stays a line break
    assert rows[2] == ["", "", "", "", ""]
    assert rows[3][0] == "-END-" and len(rows) == 4
    assert "N/A" in text                                             # agreement ticked: re-inspection / NCR show N/A


def test_instance_exports_results_and_na_reasons(chk_env):
    r = chk_env.login("chk_viewer").get(f"/api/checklist/{INSTANCE}/export-docx")
    assert r.status_code == 200, r.text
    rows = _items_rows(r)
    assert rows[0][3:] == ["Mill certs OK", "合格 Pass"]
    assert rows[1][3:] == ["Porosity at J12", "不合格 Fail"]
    assert rows[2][4] == "不適用 N/A\n不適用原因 N/A reason: Coating by others"
    assert rows[3][4] == "不適用 N/A\n歷史資料未記錄原因 No reason recorded (historical data)"
    assert rows[4][4] == "未知值 Unknown value: ??"                    # shown as stored, never reinterpreted
    text = _texts(r)
    for s in ["☐ All inspection has been done", "☑ Unfinished improvement", "NCR-007", "2026-10-20",
              "QTS-EXP1-HL-NOI-000001", "2026-10-12", "FAIL"]:
        assert s in text, s


def test_missing_checklist_is_404(chk_env):
    r = chk_env.login("chk_viewer").get("/api/checklist/no-such-id/export-docx")
    assert r.status_code == 404, r.text


def test_another_projects_checklist_is_404_like_reading_it(chk_env):
    c = chk_env.login("chk_viewer")
    assert c.get(f"/api/checklist/{OTHER_PROJECT}/").status_code == 404
    r = c.get(f"/api/checklist/{OTHER_PROJECT}/export-docx")
    assert r.status_code == 404, r.text
    assert "QTS-EXP2" not in r.text


def test_without_checklist_view_is_403(chk_env):
    r = chk_env.login("chk_noperm").get(f"/api/checklist/{TEMPLATE}/export-docx")
    assert r.status_code == 403, r.text
    assert perms.CHECKLIST_VIEW in r.json()["detail"]


def test_not_logged_in_is_refused(chk_env):
    from fastapi.testclient import TestClient
    r = TestClient(chk_env.app).get(f"/api/checklist/{TEMPLATE}/export-docx")
    assert r.status_code == 401, r.text
