"""Audit agenda (AUDIT-AGENDA-2026-001): the daily schedule rows saved with an audit — validated, sorted, stored as JSON,
covered by the Closed / Void lock, and printed in the Word export as section 4.

Real login/routes on a disposable file DB (fixture shared with test_audit_hardening_http); no main import, no development
database access.
"""
from io import BytesIO

from docx import Document

from test_audit_hardening_http import create, env, insert_raw, stored  # noqa: F401  (env is a fixture)

ROWS = [
    {'id': 3, 'date': '2026-10-02', 'start': '15:00', 'end': '15:30', 'item': 'Closing meeting', 'auditee': 'All', 'auditor': 'Wang'},
    {'id': 1, 'date': '2026-10-01', 'start': '09:00', 'end': '09:30', 'item': 'Opening meeting', 'auditee': 'All', 'auditor': 'Wang',
     'remarks': 'Room A'},
    {'id': 2, 'date': '2026-10-01', 'start': '09:30', 'end': '12:00', 'item': '8.4 External providers', 'auditee': 'Purchasing / Chen',
     'auditor': 'Wang'},
    {'id': 4, 'date': '', 'start': '', 'end': '', 'item': 'Document review (date tbc)'},
]


def agenda_rows(response):
    doc = Document(BytesIO(response.content))
    for table in doc.tables:
        if table.rows[0].cells[0].text.startswith('日期'):
            return [[c.text for c in row.cells] for row in table.rows[1:]]
    return None


def test_agenda_is_saved_sorted_and_returned(env):
    response = create(env, agenda=ROWS)
    assert response.status_code == 200, response.text
    body = response.json()
    assert [r['id'] for r in body['agenda']] == [1, 2, 3, 4]  # by date, then start; the undated row last
    assert body['agenda'][0] == {'id': 1, 'date': '2026-10-01', 'start': '09:00', 'end': '09:30', 'item': 'Opening meeting',
                                 'auditee': 'All', 'auditor': 'Wang', 'remarks': 'Room A'}
    assert env.client.get(f"/api/audit/{body['id']}").json()['agenda'] == body['agenda']
    # update replaces the list (and sorts it again)
    response = env.client.put(f"/api/audit/{body['id']}", json={'agenda': [ROWS[0], ROWS[1]]})
    assert response.status_code == 200, response.text
    assert [r['id'] for r in response.json()['agenda']] == [1, 3]


def test_an_audit_without_agenda_still_works(env):
    created = create(env).json()  # the wizard body has no agenda key at all (older frontend)
    assert created['agenda'] == []
    legacy = insert_raw(env, agenda=None)
    assert env.client.get(f'/api/audit/{legacy}').json()['agenda'] == []
    odd = insert_raw(env, agenda='not json')
    assert env.client.get(f'/api/audit/{odd}').json()['agenda'] == []  # a bad stored value must not break the read


def test_invalid_rows_are_refused(env):
    for bad in [{'date': '2026/10/01'}, {'date': '2026-02-30'}, {'start': '9:00'}, {'start': '24:00'},
                {'start': '1\u0660:00'},  # an Arabic-Indic digit: HH:MM means ASCII digits only
                {'start': '10:00', 'end': '09:00'}]:
        response = create(env, agenda=[bad])
        assert response.status_code == 422, (bad, response.text)
    created = create(env).json()
    response = env.client.put(f"/api/audit/{created['id']}", json={'agenda': [{'start': '10:00', 'end': '09:59'}]})
    assert response.status_code == 422, response.text
    # a row may have only a start, only an end, or the same start and end
    ok = create(env, agenda=[{'start': '10:00'}, {'end': '11:00'}, {'start': '12:00', 'end': '12:00'}])
    assert ok.status_code == 200, ok.text


def test_closed_audit_agenda_is_locked_but_an_unchanged_resave_passes(env):
    created = create(env, agenda=ROWS, status='Completed').json()
    assert env.client.put(f"/api/audit/{created['id']}", json={'status': 'Closed'}).status_code == 200
    # the wizard resends the whole record: the same agenda (in any order) is not a change
    same = env.client.put(f"/api/audit/{created['id']}", json={'agenda': list(reversed(ROWS)), 'status': 'Closed'})
    assert same.status_code == 200, same.text
    changed = env.client.put(f"/api/audit/{created['id']}", json={'agenda': ROWS[:2], 'status': 'Closed'})
    assert changed.status_code == 400, changed.text
    assert len(env.client.get(f"/api/audit/{created['id']}").json()['agenda']) == 4


def test_word_export_has_the_agenda_as_section_4(env):
    created = create(env, agenda=ROWS).json()
    response = env.client.get(f"/api/audit/{created['id']}/export-docx")
    assert response.status_code == 200, response.text
    doc = Document(BytesIO(response.content))
    headings = [t.rows[0].cells[0].text for t in doc.tables if len(t.columns) == 1 and t.rows[0].cells[0].text[:2] in
                {'1.', '2.', '3.', '4.', '5.', '6.', '7.'}]
    assert headings == ['1. 稽核資訊  Audit Information', '2. 稽核團隊  Audit Team', '3. 稽核範圍與準則  Scope & Criteria',
                        '4. 稽核行程  Agenda', '5. 查檢結果  Checklist Results', '6. 稽核發現與總結  Findings & Summary',
                        '7. 簽核  Sign-off']
    rows = agenda_rows(response)
    assert rows == [
        ['2026-10-01', '09:00–09:30', 'Opening meeting', 'All', 'Wang', 'Room A'],
        ['2026-10-01', '09:30–12:00', '8.4 External providers', 'Purchasing / Chen', 'Wang', '—'],
        ['2026-10-02', '15:00–15:30', 'Closing meeting', 'All', 'Wang', '—'],
        ['—', '—', 'Document review (date tbc)', '—', '—', '—'],
    ]


def test_word_export_without_agenda_and_with_odd_stored_rows(env):
    created = create(env).json()
    text = '\n'.join(p.text for p in Document(BytesIO(env.client.get(f"/api/audit/{created['id']}/export-docx").content)).paragraphs)
    assert '（尚無行程 No agenda）' in text
    odd = insert_raw(env, agenda='[1, "x", null, {"date": "2026-10-01", "start": "08:00", "item": "Only start"}]')
    response = env.client.get(f'/api/audit/{odd}/export-docx')
    assert response.status_code == 200, response.text
    assert agenda_rows(response) == [['2026-10-01', '08:00', 'Only start', '—', '—', '—']]
