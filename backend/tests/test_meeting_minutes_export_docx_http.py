"""Meeting Minutes Word export (MEETING-TABS-2026-001 R15, user 2026-10-11: "要補word匯出").

GET /api/meeting-minutes/{id}/export-docx returns the printed minutes as .docx: masthead, basic information and
attendees, the discussion register (two and three levels; progress notes in date order), attachments, sign-off —
portrait A4 — then the photo appendix (the Attachments tab's images, landscape A4, 2 × 2 per page; user decision
2026-10-11). Needs meeting:view and the record in the caller's scope (another contractor's meeting is a 404).
Same throwaway world as test_meeting_minutes_rev_http (real logins, real HTTP); uploads go to a temporary directory.
"""
import struct
import zlib
from io import BytesIO
from urllib.parse import unquote

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT

import models
from core import perms
from core.security import get_password_hash
from test_itr_revoke_approval_acceptance import PW, _get_or_create_perm
from test_meeting_minutes_rev_http import meeting_env  # noqa: F401  (fixture re-export)

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

TWO_LEVEL = [
    {"no": "1", "kind": "category", "level": 0, "content": "Technical Issues"},
    {"no": "1.1", "kind": "item", "level": 1, "content": "Safety lifeline lacks certificates\nsecond line", "owner": "TECO",
     "raisedDate": "2026-07-14", "targetDate": "2026-08-01", "status": "Open",
     "updates": [{"date": "2026-07-28", "content": "this week"},
                 {"date": "2026-07-21", "content": "earlier meeting", "carried": True},
                 {"date": "2026-07-24", "content": "filled in later"},
                 {"date": "2026-07-28", "content": "same day, written after"}]},
    {"no": "1.2", "kind": "item", "level": 1, "content": "Closed one", "owner": "HL", "status": "Closed", "updates": []},
]
THREE_LEVEL = [
    {"no": "1", "kind": "category", "level": 0, "content": "Tech"},
    {"no": "1.1", "kind": "subcategory", "level": 1, "content": "Civil OB"},
    {"no": "1.1.1", "kind": "item", "level": 2, "content": "old item", "owner": "TECO", "status": "Open"},
    {"no": "2", "topic": "Flat", "discussion": "talk", "decision": "agreed"},   # legacy flat row: a heading
]


def png(width, height) -> bytes:
    """A real (tiny, grey) PNG without Pillow."""
    raw = b"".join(b"\x00" + b"\x80" * width for _ in range(height))

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 0, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


@pytest.fixture
def uploads(tmp_path, monkeypatch):
    root = tmp_path / "uploads"
    root.mkdir()
    monkeypatch.setattr("core.uploads.upload_root", lambda: str(root))
    return root


def add_user(env, username, codes, vendor_id=None):
    with env.Session() as db:
        role = models.Role(name=f"role-{username}")
        role.permissions_rel = [_get_or_create_perm(db, c) for c in codes]
        db.add(role)
        db.flush()
        db.add(models.User(username=username, email=f"{username}@example.com", is_active=True, role_id=role.id,
                           vendor_id=vendor_id, hashed_password=get_password_hash(PW)))
        db.commit()


def attach(env, root, meeting_id, name, data, mime, at):
    (root / "meeting").mkdir(exist_ok=True)
    (root / "meeting" / name).write_bytes(data)
    with env.Session() as db:
        db.add(models.Attachment(id=f"att-{name}", entity_type="meeting", entity_id=meeting_id, file_name=name,
                                 file_path=f"meeting/{name}", mime_type=mime, category="attachment", uploaded_at=at,
                                 is_deleted=False))
        db.commit()


def create(c, **kw):
    r = c.post("/api/meeting-minutes/", json={"status": "Draft", "title": "Weekly site meeting", "vendor": "Accept Co", **kw})
    assert r.status_code == 200, r.text
    return r.json()


def texts(response):
    """Body paragraphs and every table cell's paragraphs, in document order (tables after the body)."""
    doc = Document(BytesIO(response.content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        seen = set()                          # the elements themselves (kept alive), not id(): ids get reused
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:          # a merged cell repeats in python-docx
                    continue
                seen.add(cell._tc)
                parts.extend(p.text for p in cell.paragraphs)
    return parts


def register_rows(response):
    doc = Document(BytesIO(response.content))
    for table in doc.tables:
        if table.rows[0].cells[0].text == "#":
            return [[c.text for c in row.cells] for row in table.rows[1:]]
    return None


def test_draft_export_has_the_printed_minutes_in_order(meeting_env, uploads):
    c = meeting_env.login("mtg_a")
    m = create(c, meetingDate="2026-07-28", meetingType="Weekly", meetingTime="15:00", meetingEndTime="17:15",
               location="Site office", organizer="Recorder Lin", discussionLog=TWO_LEVEL,
               attendees=[{"company": "TECO", "name": "Toby Tseng", "role": "QAQC Manager"}])
    r = c.get(f"/api/meeting-minutes/{m['id']}/export-docx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == DOCX_TYPE
    doc = Document(BytesIO(r.content))
    assert doc.sections[0].orientation == WD_ORIENT.PORTRAIT
    assert round(doc.sections[0].page_width.cm, 1) == 21.0 and round(doc.sections[0].page_height.cm, 1) == 29.7
    assert len(doc.sections) == 1                                        # no photos: no appendix
    all_text = "\n".join(texts(r))
    for s in ["會議記錄 MEETING MINUTES", m["documentNumber"], "版次 Rev：1.0", "DRAFT", "Weekly site meeting",
              "28 Jul 2026　Weekly", "Accept Co", "Site office　15:00 ~ 17:15", "Recorder Lin",
              "Toby Tseng", "QAQC Manager", "（無附件 No attachments）", "4.1 記錄人 Recorded By"]:
        assert s in all_text, s
    rows = register_rows(r)
    assert [row[0] for row in rows] == ["1.0", "1.1", "1.2"]
    assert rows[0][1] == rows[0][3] == "Technical Issues"                # heading spans the row
    assert rows[1][2] == "TECO" and rows[1][3] == "1 Aug 2026\nOpen" and rows[2][3] == "Closed"
    # raised date, text, then the notes by date (the backfilled day in its place, same day as written)
    assert rows[1][1].split("\n") == ["14 Jul 2026", "Safety lifeline lacks certificates", "second line",
                                      "21 Jul 2026", "earlier meeting", "24 Jul 2026", "filled in later",
                                      "28 Jul 2026", "this week", "28 Jul 2026", "same day, written after"]


def test_three_level_and_legacy_rows_and_the_published_record(meeting_env, uploads):
    c = meeting_env.login("mtg_a")
    m = create(c, discussionLog=THREE_LEVEL)
    assert c.put(f"/api/meeting-minutes/{m['id']}", json={"status": "Published"}).status_code == 200
    r = c.get(f"/api/meeting-minutes/{m['id']}/export-docx")
    assert r.status_code == 200, r.text
    assert "PUBLISHED" in "\n".join(texts(r))
    rows = register_rows(r)
    assert [(row[0], row[1]) for row in rows] == [("1.0", "Tech"), ("1.1", "Civil OB"), ("1.1.1", "old item"),
                                                  ("2.0", "Flat / talk / agreed")]


def test_filename_carries_number_and_rev(meeting_env, uploads):
    c = meeting_env.login("mtg_a")
    first = create(c)
    second = c.post(f"/api/meeting-minutes/{first['id']}/new-occurrence").json()
    disposition = c.get(f"/api/meeting-minutes/{second['id']}/export-docx").headers["content-disposition"]
    assert f"filename*=UTF-8''{second['documentNumber']}_Rev2.0.docx" in unquote(disposition)
    assert 'filename="' in disposition and disposition.split('filename="')[1].split('"')[0].endswith("_Rev2_0.docx")


def test_photo_appendix_landscape_four_per_page(meeting_env, uploads):
    c = meeting_env.login("mtg_a")
    m = create(c)
    for i in range(5):     # five photos: two pages (4 + 1); a tall one must fit its box too
        attach(meeting_env, uploads, m["id"], f"photo{i + 1}.png", png(40, 90) if i == 1 else png(80, 60), "image/png", f"2026-10-11T00:0{i}:00")
    attach(meeting_env, uploads, m["id"], "drawing.pdf", b"%PDF-1.4 not a photo", "application/pdf", "2026-10-11T00:09:00")
    attach(meeting_env, uploads, m["id"], "broken.png", b"not really a png", "image/png", "2026-10-11T00:08:00")
    r = c.get(f"/api/meeting-minutes/{m['id']}/export-docx")
    assert r.status_code == 200, r.text
    doc = Document(BytesIO(r.content))
    assert [s.orientation for s in doc.sections] == [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE]
    assert round(doc.sections[1].page_width.cm, 1) == 29.7
    assert len(doc.inline_shapes) == 5                                   # the PDF and the broken image are not embedded
    for shape in doc.inline_shapes:
        assert shape.width.cm <= 12.41 and shape.height.cm <= 7.01
    all_text = "\n".join(texts(r))
    assert "1. photo1.png　（照片 1，見照片附錄 See photo appendix）" in all_text
    assert "7. drawing.pdf" in all_text and "drawing.pdf　（照片" not in all_text
    assert "6. broken.png　（照片 6，" in all_text
    assert f"照片 Photos — {m['documentNumber']} (1.0)　　1 / 2" in all_text and "2 / 2" in all_text
    assert "照片 Photo 5　photo5.png" in all_text
    assert "（圖片無法顯示 image could not be shown）" in all_text      # the broken one: a note, not a failed export


def test_a_file_path_outside_the_uploads_is_not_read(meeting_env, uploads, tmp_path):
    secret = tmp_path / "secret.png"
    secret.write_bytes(png(10, 10))
    c = meeting_env.login("mtg_a")
    m = create(c)
    with meeting_env.Session() as db:
        db.add(models.Attachment(id="att-escape", entity_type="meeting", entity_id=m["id"], file_name="escape.png",
                                 file_path="../secret.png", mime_type="image/png", category="attachment",
                                 uploaded_at="2026-10-11T00:00:00", is_deleted=False))
        db.commit()
    r = c.get(f"/api/meeting-minutes/{m['id']}/export-docx")
    assert r.status_code == 200
    assert len(Document(BytesIO(r.content)).inline_shapes) == 0


def test_permission_and_scope(meeting_env, uploads):
    c = meeting_env.login("mtg_a")
    m = create(c)
    with meeting_env.Session() as db:
        db.add(models.Contractor(id="MTG-V2", name="Other Co", abbreviation="OTH"))
        db.commit()
    add_user(meeting_env, "mtg_noview", [perms.MEETING_CREATE])
    add_user(meeting_env, "mtg_other", [perms.MEETING_VIEW], vendor_id="MTG-V2")
    add_user(meeting_env, "mtg_same", [perms.MEETING_VIEW], vendor_id="MTG-V1")
    assert meeting_env.login("mtg_noview").get(f"/api/meeting-minutes/{m['id']}/export-docx").status_code == 403
    assert meeting_env.login("mtg_other").get(f"/api/meeting-minutes/{m['id']}/export-docx").status_code == 404
    assert meeting_env.login("mtg_same").get(f"/api/meeting-minutes/{m['id']}/export-docx").status_code == 200
    assert c.get("/api/meeting-minutes/does-not-exist/export-docx").status_code == 404


def test_odd_values_in_the_free_json_do_not_break_the_export(meeting_env, uploads):
    """discussionLog / attendees are free JSON: numbers where text is expected are printed as text, not a 500."""
    c = meeting_env.login("mtg_a")
    m = create(c, attendees=[{"company": 7, "name": None, "role": ["x"]}, "not a dict"], discussionLog=[
        {"no": "1", "kind": "category", "content": 2026},
        {"no": "1.1", "kind": "item", "content": 42, "owner": 3, "targetDate": 20260801, "status": "Open",
         "updates": [{"date": 20260720, "content": "numeric day"}, {"date": "2026-07-19", "content": "text day"}, {"content": 5}]},
    ])
    r = c.get(f"/api/meeting-minutes/{m['id']}/export-docx")
    assert r.status_code == 200, r.text
    rows = register_rows(r)
    assert rows[0][1] == "2026" and rows[1][1].startswith("42") and rows[1][2] == "3" and rows[1][3].startswith("20260801")
