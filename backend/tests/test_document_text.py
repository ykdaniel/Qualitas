"""core.document_text — Word text extraction for the controlled-document search (KM-DOCS-2026-001 D1, Q4)."""
import io

from docx import Document

from core.document_text import extract_docx_text, snippet


def _docx(paragraphs=(), table=None, header=None, footer=None) -> bytes:
    doc = Document()
    if header:
        doc.sections[0].header.paragraphs[0].text = header
    if footer:
        doc.sections[0].footer.paragraphs[0].text = footer
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_paragraphs_tables_header_and_footer_are_extracted_in_order():
    text = extract_docx_text(_docx(paragraphs=["焊接前檢查   程序", "Scope"], table=[["項目", "Item"], ["數量", "Qty"]],
                                   header="QTS-KM-000001 Rev1.0", footer="Page"))
    lines = text.split("\n")
    assert lines[0] == "QTS-KM-000001 Rev1.0"             # header first
    assert "焊接前檢查 程序" in lines                      # whitespace runs collapse
    assert "項目 | Item" in lines and "數量 | Qty" in lines
    assert lines[-1] == "Page"                            # footer last
    assert lines.index("Scope") < lines.index("項目 | Item")


def test_not_a_docx_yields_none_and_empty_yields_none():
    assert extract_docx_text(b"%PDF-1.4 not a word file") is None
    assert extract_docx_text(b"") is None


def test_snippet_gives_context_around_the_first_case_insensitive_hit():
    text = "a" * 200 + "Welding Procedure" + "b" * 200
    s = snippet(text, "procedure", width=10)
    assert s == "…" + "aa" + "Welding Procedure" + "b" * 10 + "…"      # 10 chars before the hit: "aa" + "Welding "
    assert snippet(text, "missing") is None
    assert snippet(None, "x") is None
    assert snippet("short Hit", "hit") == "short Hit"
